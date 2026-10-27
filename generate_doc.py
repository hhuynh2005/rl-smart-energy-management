import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Set standard academic margins: 2cm top/bottom, 2.5cm left, 2cm right
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.8)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(30, 30, 30)

    # --- TITLE / HEADER ---
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_inst.add_run("BÁO CÁO THỰC HÀNH 1: HỌC TĂNG CƯỜNG (REINFORCEMENT LEARNING)\nCHUẨN ĐÁNH GIÁ 4 YÊU CẦU CỐT LÕI TẠI BUỔI THỰC HÀNH")
    run_inst.bold = True
    run_inst.font.size = Pt(12.5)
    run_inst.font.color.rgb = RGBColor(70, 70, 70)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("ĐỀ TÀI: ĐIỀU KHIỂN NĂNG LƯỢNG THÔNG MINH\n(SMART ENERGY MANAGEMENT SYSTEM - SEMS-RL)")
    run_title.bold = True
    run_title.font.size = Pt(16)
    run_title.font.color.rgb = RGBColor(16, 78, 139)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Ứng dụng Học tăng cường trong Quản lý Tối ưu Pin lưu trữ BESS và Điện mặt trời Solar PV\nThời gian báo cáo: Sáng Thứ 7, ngày 03/10/2026")
    run_sub.italic = True
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(100, 100, 100)

    # Info Box Table
    table_info = doc.add_table(rows=7, cols=2)
    table_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_info.autofit = False

    all_info = [
        ("Nội dung đánh giá TH1:", "Bao gồm đầy đủ 4 yêu cầu từ Giảng viên (Đề xuất đề tài, Tính năng phần mềm, Phân công công việc, Mô hình hóa MDP)"),
        ("Cơ cấu nhân sự (5 người):", "Phân chia chuyên sâu: 2 thành viên Môi trường + 3 thành viên Thuật toán"),
        ("1. Nguyễn Văn Huỳnh (Nhóm trưởng)", "MSSV: .................... - Phụ trách: MÔI TRƯỜNG (Kiến trúc & Động lực học pin)"),
        ("2. Lý Đình Sơn", "MSSV: .................... - Phụ trách: MÔI TRƯỜNG (Xử lý dữ liệu chuỗi thời gian & Tích hợp Gym Env)"),
        ("3. Nguyễn Văn Trường", "MSSV: .................... - Phụ trách: THUẬT TOÁN (Lõi chính: Soft Actor-Critic - SAC Lead)"),
        ("4. Phạm Quang Anh", "MSSV: .................... - Phụ trách: THUẬT TOÁN (Mô hình so sánh: TD3 & Rule-based Baseline)"),
        ("5. Nguyễn Tùng Dương", "MSSV: .................... - Phụ trách: THUẬT TOÁN (Mô hình so sánh: PPO & Đánh giá so chuẩn)"),
    ]

    for idx, (col1, col2) in enumerate(all_info):
        row = table_info.rows[idx]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(2.8)
        c2.width = Inches(4.0)
        set_cell_background(c1, "F2F5F8")
        set_cell_background(c2, "FAFBFC")
        set_cell_margins(c1, 70, 70, 120, 120)
        set_cell_margins(c2, 70, 70, 120, 120)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(col1)
        r1.bold = True
        r1.font.size = Pt(10.5)
        r1.font.name = 'Times New Roman'
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(col2)
        r2.font.size = Pt(10.5)
        r2.font.name = 'Times New Roman'
        if idx == 0:
            r2.bold = True
            r2.font.color.rgb = RGBColor(180, 40, 40)
        elif idx == 1:
            r2.bold = True
            r2.font.color.rgb = RGBColor(16, 78, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Function helpers
    def add_sec_heading(title_text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(title_text)
        run.bold = True
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(16, 78, 139)
        return h

    def add_sub_heading(sub_text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(sub_text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(34, 102, 102)
        return h

    def add_body(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
        rt = p.add_run(text)
        rt.italic = italic
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
        p.add_run(text)
        return p

    # =========================================================================
    # YÊU CẦU 1: ĐỀ XUẤT ĐỀ TÀI / DỰ ÁN PHẦN MỀM
    # =========================================================================
    add_sec_heading("YÊU CẦU 1: ĐỀ XUẤT ĐỀ TÀI / DỰ ÁN PHẦN MỀM CÓ SỬ DỤNG HỌC TĂNG CƯỜNG")
    
    add_sub_heading("1.1. Tên đề tài và Mô tả bài toán")
    add_body("• Tên đề tài: Hệ thống Quản lý và Điều khiển Năng lượng Thông minh bằng Học tăng cường (Smart Energy Management System with Reinforcement Learning - SEMS-RL).")
    add_body("• Mô tả bài toán: Xây dựng một tác tử thông minh (RL Agent) có khả năng tự động điều phối, tối ưu hóa việc nạp/xả của hệ thống pin lưu trữ năng lượng (Battery Energy Storage System - BESS) kết hợp cùng hệ thống điện mặt trời áp mái (Solar PV) tại các nhà máy, tòa nhà thông minh hoặc lưới điện siêu nhỏ (Microgrid).")
    add_body("• Ứng dụng thực tế: Nhà máy sản xuất công nghiệp, tòa nhà văn phòng thông minh (Smart Buildings), trang trại điện mặt trời kết hợp lưu trữ BESS, và các công ty năng lượng cần cân bằng tải.")

    add_sub_heading("1.2. Động lực và Lý do lựa chọn Học tăng cường (RL)")
    add_body("Trong thực tế vận hành hệ thống năng lượng, các phương pháp truyền thống bộc lộ nhiều điểm yếu chí mạng:")
    add_bullet(" Biểu giá điện của EVN chênh lệch cực lớn giữa giờ cao điểm (giá rất cao), giờ bình thường và giờ thấp điểm (giá rẻ). Nếu tiêu thụ nhiều điện lưới vào giờ cao điểm, chi phí sẽ tăng đột biến.", "1. Biến động biểu giá điện theo giờ (TOU Tariff):")
    add_bullet(" Bức xạ mặt trời phụ thuộc thời tiết, mây che, mang tính ngẫu nhiên và bấp bênh. Đỉnh phát quang điện thường rơi vào giữa trưa (11h-13h) trong khi đỉnh phụ tải tiêu thụ thường là đầu giờ sáng hoặc chiều tối.", "2. Tính bất định cao của năng lượng tái tạo:")
    add_bullet(" Nhu cầu sử dụng điện đột biến làm phát sinh chi phí phạt công suất đỉnh (Demand Charge) và nguy cơ quá tải trạm biến áp.", "3. Rủi ro quá tải đỉnh phụ tải (Peak Load):")
    add_bullet(" Các thuật toán Heuristic dựa trên luật cố định (Rule-based: cứ có nắng là nạp, hết nắng xả) rất cứng nhắc; còn thuật toán Quy hoạch tuyến tính (Linear Programming - LP) yêu cầu phải biết trước hoàn hảo thông tin tương lai — điều bất khả thi.", "4. Hạn chế của giải pháp truyền thống:")
    add_body("👉 Giải pháp với Học tăng cường (RL): Tác tử RL học được chính sách hành động thông qua việc tương tác liên tục với môi trường mô phỏng. RL có khả năng dự phòng chiến lược dài hạn (ví dụ: chủ động sạc pin lúc 2h sáng khi điện rẻ nhất để chuẩn bị xả vào giờ cao điểm 9h sáng và 18h tối), tự động thích ứng với biến động ngẫu nhiên của thời tiết và tải mà không cần giả định môi trường tĩnh.")

    add_sub_heading("1.3. Các mục tiêu chính của dự án")
    add_bullet(" Giảm thiểu tối đa chi phí mua điện từ lưới điện quốc gia nhờ cơ chế dịch chuyển phụ tải (Load Shifting) và cắt đỉnh phụ tải (Peak Shaving).", "• Mục tiêu 1 - Tối ưu hóa chi phí điện năng (Cost Minimization):")
    add_bullet(" Tận dụng tối đa nguồn điện sạch tại chỗ, giảm thiểu việc phát ngược điện thừa lên lưới với giá rẻ hoặc lãng phí công suất mặt trời.", "• Mục tiêu 2 - Tối đa hóa tỷ lệ tự tiêu thụ năng lượng xanh (Self-Consumption):")
    add_bullet(" Kiểm soát tốc độ sạc xả, ngăn chặn hiện tượng sạc xả liên tục đổi chiều, giữ mức pin trong dải an toàn (20% - 90%) nhằm kéo dài tuổi thọ cell pin.", "• Mục tiêu 3 - Bảo vệ tuổi thọ hệ thống pin BESS (Battery Degradation):")
    add_bullet(" Đóng gói thành phần mềm hoàn chỉnh có API dịch vụ, bảng điều khiển Dashboard và bộ lọc an toàn vật lý.", "• Mục tiêu 4 - Xây dựng giải pháp phần mềm thông minh:")

    add_sub_heading("1.4. Công nghệ và Nền tảng dự kiến sử dụng")
    add_bullet(" Nhóm triển khai 3 thuật toán Continuous RL: Soft Actor-Critic (SAC - Thuật toán lõi chính, Off-policy, Maximum Entropy), TD3 (Twin Delayed DDPG - So sánh, chống overestimation) và PPO (Proximal Policy Optimization - So sánh, On-policy chuẩn mực).", "• Thuật toán Học tăng cường:")
    add_bullet(" Python 3.10+, Gymnasium (chuẩn OpenAI Gym), PyTorch, Stable-Baselines3.", "• Thư viện & Framework:")
    add_bullet(" Dữ liệu phụ tải và bức xạ mặt trời thực tế từ Pecan Street Project / CityLearn / NREL, kết hợp biểu giá điện TOU theo quy định của EVN.", "• Tập dữ liệu (Datasets):")
    add_bullet(" FastAPI (REST API phục vụ inference thời gian thực) và Streamlit / ReactJS (Dashboard trực quan hóa).", "• Nền tảng ứng dụng:")

    # =========================================================================
    # YÊU CẦU 2: ĐỀ XUẤT CÁC TÍNH NĂNG CỦA PHẦN MỀM
    # =========================================================================
    add_sec_heading("YÊU CẦU 2: ĐỀ XUẤT CÁC TÍNH NĂNG CỦA PHẦN MỀM (SEMS-RL)")
    add_body("Hệ thống phần mềm SEMS-RL được thiết kế thành một nền tảng quản lý năng lượng toàn diện, gồm 6 phân hệ tính năng cốt lõi:")

    table_features = doc.add_table(rows=1, cols=3)
    table_features.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_features.autofit = False

    hdr = table_features.rows[0].cells
    hdr[0].width = Inches(0.8)
    hdr[1].width = Inches(2.2)
    hdr[2].width = Inches(3.8)
    
    headers = ["STT", "Tính năng / Phân hệ", "Mô tả chi tiết chức năng"]
    for i, h_text in enumerate(headers):
        set_cell_background(hdr[i], "1F4E79")
        set_cell_margins(hdr[i], 90, 90, 100, 100)
        p = hdr[i].paragraphs[0]
        if i == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10.5)

    features_data = [
        ("1", "Môi trường Năng lượng Số\n(Energy Digital Twin)", "Mô phỏng chính xác hệ thống pin BESS (dung lượng kWh, hiệu suất nạp xả, tổn thất nhiệt eta, suy thoái pin), tích hợp nguồn phát điện mặt trời (PV), biểu giá điện TOU và phụ tải tòa nhà theo từng bước 15 phút."),
        ("2", "Bộ điều khiển RL Tự động\n(Autonomous RL Controller)", "Tích hợp mô hình RL liên tục (SAC/TD3) đưa ra quyết định công suất sạc/xả tối ưu theo thời gian thực (real-time control) nhằm tối thiểu hóa chi phí và cân bằng tải."),
        ("3", "Module Đối chuẩn Hiệu năng\n(Benchmarking Engine)", "Tự động so sánh hiệu quả giữa Agent RL với các phương pháp đối chuẩn: Rule-based Heuristic (luật cố định theo giờ), Không có pin lưu trữ, và Quy hoạch tuyến tính lý tưởng (Linear Programming)."),
        ("4", "Bảng điều khiển Giám sát\n(Real-time Dashboard UI)", "Giao diện trực quan hóa dòng luân chuyển năng lượng (PV -> Battery -> Load/Grid), biểu đồ % pin SoC, đồ thị phụ tải san phẳng (Peak Shaving) và số tiền tiết kiệm lũy kế."),
        ("5", "Bộ lọc An toàn Phần cứng\n(Safety Guardrails)", "Cơ chế bảo vệ vật lý: tự động chặn các hành động sạc/xả vượt quá công suất định mức; tự động ngắt khi pin chạm ngưỡng cạn (20%) hoặc ngưỡng đầy (90%)."),
        ("6", "Chế độ Can thiệp Thủ công\n(Manual Override Mode)", "Cho phép kỹ sư vận hành chuyển đổi linh hoạt giữa chế độ điều khiển tự động hoàn toàn bằng RL và chế độ điều khiển thủ công theo chỉ định khi xảy ra bảo trì hoặc tình huống khẩn cấp.")
    ]

    for item in features_data:
        row = table_features.add_row()
        for c_idx, text in enumerate(item):
            cell = row.cells[c_idx]
            cell.width = hdr[c_idx].width
            set_cell_margins(cell, 70, 70, 90, 90)
            if c_idx % 2 == 1:
                set_cell_background(cell, "F7F9FB")
            p = cell.paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # YÊU CẦU 3: PHÂN CÔNG CÔNG VIỆC DỰ KIẾN
    # =========================================================================
    add_sec_heading("YÊU CẦU 3: PHÂN CÔNG CÔNG VIỆC DỰ KIẾN CHO CÁC THÀNH VIÊN")
    add_body("Để phát huy tối đa năng lực và chuyên môn hóa, nhóm 5 thành viên được phân chia rõ ràng thành 2 phân hệ chuyên trách: 2 thành viên phụ trách Môi trường & Dữ liệu, và 3 thành viên phụ trách Thuật toán & Đối chuẩn:")

    table_members = doc.add_table(rows=1, cols=4)
    table_members.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_members.autofit = False

    m_hdr = table_members.rows[0].cells
    m_hdr[0].width = Inches(1.5)
    m_hdr[1].width = Inches(1.4)
    m_hdr[2].width = Inches(2.6)
    m_hdr[3].width = Inches(1.5)

    m_headers = ["Họ và tên", "Phân nhóm & Vai trò", "Nhiệm vụ chính phụ trách", "Sản phẩm đầu ra"]
    for i, h_text in enumerate(m_headers):
        set_cell_background(m_hdr[i], "1F4E79")
        set_cell_margins(m_hdr[i], 90, 90, 80, 80)
        p = m_hdr[i].paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)

    members_data = [
        ("1. Nguyễn Văn Huỳnh\n(Nhóm trưởng)", 
         "NHÓM MÔI TRƯỜNG\n(Architecture & Lead)", 
         "- Quản lý tiến độ chung của toàn nhóm.\n- Thiết kế kiến trúc tổng thể môi trường mô phỏng theo chuẩn Gymnasium (`MicrogridBessEnv`).\n- Xây dựng mô hình toán động lực học pin BESS (phương trình nạp/xả, tổn thất nhiệt eta, suy thoái pin theo chu kỳ sạc/xả).\n- Định nghĩa không gian trạng thái S, không gian hành động A và hàm chuyển trạng thái P.",
         "- Module `MicrogridBessEnv`.\n- Bản đặc tả mô hình toán học MDP.\n- Phương trình động lực học pin."),
        
        ("2. Lý Đình Sơn", 
         "NHÓM MÔI TRƯỜNG\n(Data & Env Integration)", 
         "- Thu thập, làm sạch và chuẩn hóa dữ liệu chuỗi thời gian: Bức xạ mặt trời (PV), phụ tải (Load), biểu giá điện TOU của EVN.\n- Lập trình Data Pipeline nạp dữ liệu vào môi trường Gym.\n- Thiết kế cơ chế an toàn (Safety Guardrails) trong môi trường: ngăn sạc/xả vượt ngưỡng, chặn lỗi vật lý khi pin cạn/đầy.\n- Tạo môi trường nhiễu ngẫu nhiên để thử nghiệm độ bền (robustness) của môi trường.",
         "- Bộ dữ liệu sạch `data/processed.csv`.\n- Pipeline nạp dữ liệu thời gian thực.\n- Module Safety Guardrails cho môi trường."),

        ("3. Nguyễn Văn Trường", 
         "NHÓM THUẬT TOÁN\n(Algorithm 1: SAC Lead)", 
         "- Chịu trách nhiệm chính về thuật toán lõi **SAC (Soft Actor-Critic)**.\n- Thiết kế mạng nơ-ron Actor-Critic (Gaussian Policy cho continuous action).\n- Lập trình cơ chế Replay Buffer, tối đa hóa Entropy tự động (Auto-tuned Temperature alpha).\n- Thiết kế chi tiết và tinh chỉnh hàm phần thưởng đa mục tiêu R (tiền điện, chai pin, cắt đỉnh tải).\n- Tinh chỉnh siêu tham số và tối ưu hóa tốc độ hội tụ của SAC.",
         "- Mã nguồn thuật toán SAC hoàn chỉnh.\n- Log huấn luyện và biểu đồ hội tụ Reward của SAC.\n- Báo cáo phân tích hàm thưởng."),

        ("4. Phạm Quang Anh", 
         "NHÓM THUẬT TOÁN\n(Algorithm 2: TD3 & Baseline)", 
         "- Chịu trách nhiệm triển khai thuật toán đối chuẩn **TD3 (Twin Delayed DDPG)**.\n- Lập trình mạng Twin Critic để khắc phục lỗi ước lượng quá cao Q-value (overestimation bias) và cập nhật Policy trễ.\n- Xây dựng thuật toán đối chuẩn thực tế **Rule-based Heuristic** (luật cố định theo khung giờ TOU) làm mốc so sánh.\n- Tinh chỉnh siêu tham số cho TD3 và so sánh đối đầu trực tiếp giữa TD3 và SAC.",
         "- Mã nguồn thuật toán TD3.\n- Mã nguồn thuật toán Rule-based Baseline.\n- Báo cáo đối chuẩn SAC vs TD3."),

        ("5. Nguyễn Tùng Dương", 
         "NHÓM THUẬT TOÁN\n(Algorithm 3: PPO & Evaluation)", 
         "- Chịu trách nhiệm triển khai thuật toán **PPO (Proximal Policy Optimization)** cho continuous control.\n- Xây dựng pipeline kiểm thử, đánh giá tổng hợp cả 3 thuật toán RL (SAC vs TD3 vs PPO) và Baseline.\n- Đo lường và trực quan hóa các chỉ số kỹ thuật: % chi phí tiết kiệm được, % cắt giảm phụ tải đỉnh (Peak Shaving Ratio).\n- Xây dựng Dashboard hiển thị kết quả và biên soạn Slide thuyết trình cho các buổi báo cáo.",
         "- Mã nguồn thuật toán PPO.\n- Pipeline đánh giá & bảng so sánh tổng hợp.\n- Slide thuyết trình & Dashboard kết quả.")
    ]

    for item in members_data:
        row = table_members.add_row()
        for c_idx, text in enumerate(item):
            cell = row.cells[c_idx]
            cell.width = m_hdr[c_idx].width
            set_cell_margins(cell, 70, 70, 80, 80)
            if c_idx == 0:
                set_cell_background(cell, "F2F5F8")
            elif c_idx == 1:
                if "MÔI TRƯỜNG" in text:
                    set_cell_background(cell, "EBF5FB")
                else:
                    set_cell_background(cell, "FEF9E7")
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # YÊU CẦU 4: MÔ HÌNH HÓA BÀI TOÁN THÀNH QUY TRÌNH QUYẾT ĐỊNH MARKOV (MDP)
    # =========================================================================
    add_sec_heading("YÊU CẦU 4: MÔ HÌNH HÓA BÀI TOÁN THÀNH QUY TRÌNH QUYẾT ĐỊNH MARKOV (MDP)")
    add_body("Hệ thống điều khiển năng lượng được mô hình hóa toán học chặt chẽ dưới dạng Quy trình quyết định Markov (Markov Decision Process - MDP), định nghĩa bởi bộ 5 thành phần:")
    
    p_mdp = doc.add_paragraph()
    p_mdp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_mdp = p_mdp.add_run("M = < S,  A,  P,  R,  γ >")
    r_mdp.bold = True
    r_mdp.font.size = Pt(13)
    r_mdp.font.color.rgb = RGBColor(16, 78, 139)

    add_body("Thời gian vận hành được rời rạc hóa thành các bước thời gian delta_t = 15 phút (0.25 giờ), tương ứng T = 96 bước quyết định cho một chu kỳ 24 giờ hoàn chỉnh.")

    add_sub_heading("4.1. Không gian Trạng thái (State Space - S)")
    add_body("Trạng thái s_t tại thời điểm t cung cấp đầy đủ thông tin về tình trạng nội tại của hệ thống pin lưu trữ và các yếu tố môi trường ngoại cảnh:")
    
    p_state = doc.add_paragraph()
    p_state.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_state = p_state.add_run("s_t = [ SoC_t,  P_load,t,  P_pv,t,  C_grid,t,  sin(2π * h_t / 24),  cos(2π * h_t / 24),  d_t ]^T")
    r_state.bold = True
    r_state.font.size = Pt(11)
    r_state.font.color.rgb = RGBColor(34, 102, 102)

    add_bullet(" Mức dung lượng hiện tại của hệ thống pin BESS (0.0 = cạn kiệt, 1.0 = đầy 100%).", "• SoC_t in [0.0, 1.0] (State of Charge):")
    add_bullet(" Tổng công suất điện mà tòa nhà/nhà máy đang tiêu thụ tại thời điểm t (đơn vị: kW).", "• P_load,t >= 0 (Building Load Demand):")
    add_bullet(" Công suất điện mặt trời phát ra tức thời tại thời điểm t (đơn vị: kW), phản ánh điều kiện bức xạ và thời tiết.", "• P_pv,t >= 0 (Solar PV Generation):")
    add_bullet(" Đơn giá mua điện từ lưới điện quốc gia tại thời điểm t theo biểu giá TOU của EVN (đơn vị: VNĐ/kWh).", "• C_grid,t > 0 (Grid Electricity Price):")
    add_bullet(" Giờ trong ngày h_t in [0, 23] được mã hóa qua hàm sin/cos chu kỳ vòng tròn, giúp mạng nơ-ron nhận thức được tính liên tục ngày/đêm và không bị gián đoạn giữa 23h và 0h.", "• sin(2π * h_t / 24) và cos(2π * h_t / 24) (Cyclical Time Encoding):")
    add_bullet(" Cờ nhị phân phân biệt ngày làm việc trong tuần (weekday = 1) và ngày nghỉ cuối tuần (weekend = 0).", "• d_t in {0, 1} (Day-of-Week Indicator):")

    add_sub_heading("4.2. Không gian Hành động (Action Space - A)")
    add_body("Do pin BESS cho phép điều chỉnh công suất sạc và xả liên tục, nhóm thiết lập không gian hành động liên tục chuẩn hóa:")
    
    p_act = doc.add_paragraph()
    p_act.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_act = p_act.add_run("a_t in [-1.0, 1.0]")
    r_act.bold = True
    r_act.font.size = Pt(12)
    r_act.font.color.rgb = RGBColor(180, 40, 40)

    add_body("Hành động chuẩn hóa a_t được chuyển đổi thành công suất sạc/xả vật lý P_batt,t như sau:")
    add_bullet(" Hệ thống tiến hành sạc điện vào pin với công suất: P_batt,t = a_t * P_charge_max", "• Khi a_t > 0 (Chế độ Sạc - Charging):")
    add_bullet(" Hệ thống pin ở trạng thái nghỉ, không sạc không xả: P_batt,t = 0", "• Khi a_t = 0 (Chế độ Nghỉ - Idle):")
    add_bullet(" Hệ thống tiến hành xả điện từ pin để cấp cho tải: P_batt,t = |a_t| * (-P_discharge_max)", "• Khi a_t < 0 (Chế độ Xả - Discharging):")

    add_sub_heading("4.3. Động lực học Chuyển trạng thái (State Transition Dynamics - P)")
    add_body("1) Phương trình cân bằng công suất tại điểm hòa lưới của tòa nhà:")
    p_pgrid = doc.add_paragraph()
    p_pgrid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pgrid.add_run("P_grid,t = P_load,t - P_pv,t + P_batt,t").bold = True
    add_bullet(" Hệ thống bị thiếu điện, phải mua điện từ lưới điện quốc gia: P_grid,buy,t = P_grid,t", "• Nếu P_grid,t > 0:")
    add_bullet(" Năng lượng mặt trời dư thừa, có thể phát ngược bán lên lưới điện: P_grid,sell,t = |P_grid,t|", "• Nếu P_grid,t < 0:")

    add_body("2) Phương trình cập nhật trạng thái pin (SoC) theo thời gian:")
    p_soc = doc.add_paragraph()
    p_soc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_soc.add_run("SoC_t+1 = SoC_t + ΔSoC_t").bold = True
    add_body("Trong đó độ thay đổi dung lượng ΔSoC_t có tính đến hiệu suất tổn hao nhiệt của pin:")
    add_bullet(" ΔSoC_t = (eta_charge * P_batt,t * delta_t) / E_capacity   (với eta_charge ≈ 0.92)", "• Khi sạc (P_batt,t > 0):")
    add_bullet(" ΔSoC_t = (P_batt,t * delta_t) / (eta_discharge * E_capacity)  (với eta_discharge ≈ 0.92)", "• Khi xả (P_batt,t < 0):")
    add_body("Với E_capacity là tổng dung lượng danh định của hệ thống pin (đơn vị: kWh).")

    add_sub_heading("4.4. Thiết kế Hàm Phần thưởng Đa mục tiêu (Reward Function - R)")
    add_body("Hàm phần thưởng r_t được thiết kế dưới dạng hàm phạt âm (Cost Penalty) cần cực tiểu hóa:")
    
    p_rew = doc.add_paragraph()
    p_rew.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_rew = p_rew.add_run("r_t = - [ C_electricity,t  +  w1 * C_degradation,t  +  w2 * C_peak,t  +  w3 * P_violation,t ]")
    r_rew.bold = True
    r_rew.font.size = Pt(11)
    r_rew.font.color.rgb = RGBColor(16, 78, 139)

    add_bullet(" C_electricity,t = (C_grid,t * P_grid,buy,t - C_feedin,t * P_grid,sell,t) * delta_t. Giúp Agent sạc lúc giá thấp hoặc thừa điện mặt trời, xả lúc giá cao để giảm hóa đơn tiền điện.", "1. Chi phí mua bán điện lưới (C_electricity,t):")
    add_bullet(" C_degradation,t = c_wear * |P_batt,t| * delta_t. Phạt sự hao mòn vật lý, hạn chế việc Agent đổi chiều sạc/xả liên tục gây chai pin.", "2. Chi phí suy thoái và chai pin (C_degradation,t):")
    add_bullet(" C_peak,t = max(0, P_grid,buy,t - P_threshold)^2. Phạt lũy tiến bậc hai khi lượng điện lấy từ lưới vượt quá công suất trần P_threshold (Peak Shaving).", "3. Phạt công suất đỉnh (C_peak,t):")
    add_bullet(" P_violation,t = (SoC_min - SoC_t+1)^2 nếu SoC_t+1 < SoC_min, hoặc (SoC_t+1 - SoC_max)^2 nếu SoC_t+1 > SoC_max. Phạt nặng khi pin vượt ngưỡng an toàn [20%, 90%].", "4. Phạt vi phạm giới hạn an toàn pin (P_violation,t):")

    add_sub_heading("4.5. Hệ số Chiết khấu và Chân trời thời gian (Discount Factor & Horizon)")
    add_bullet(" Hệ số chiết khấu γ = 0.99: Bảo đảm Agent có tầm nhìn dài hạn, chủ động sạc lúc 2h sáng để xả vào lúc 18h tối.", "• Hệ số chiết khấu γ = 0.99:")
    add_bullet(" Chu kỳ 24 giờ gồm T = 96 bước thời gian liên tiếp (delta_t = 15 phút).", "• Chân trời thời gian (Horizon T = 96 steps):")

    # =========================================================================
    # KẾ HOẠCH TRIỂN KHAI TIẾP THEO
    # =========================================================================
    add_sec_heading("V. KẾ HOẠCH BÁO CÁO VÀ KỲ VỌNG ĐẠT ĐƯỢC")
    add_bullet(" Báo cáo Slide thuyết trình (10 slide) bảo vệ 4 nội dung đánh giá trước Giảng viên và lớp.", "• Buổi thực hành 1 (Sáng thứ 7, 03/10/2026):")
    add_bullet(" Hoàn thiện môi trường Gym và kiểm thử dữ liệu chuỗi thời gian (Huỳnh & Sơn).", "• Buổi thực hành 2 & 3:")
    add_bullet(" Huấn luyện SAC (Trường), TD3 & Rule-based (Quang Anh), PPO & Benchmark (Dương).", "• Buổi thực hành 4 & 5:")
    add_bullet(" Tích hợp mô hình vào Dashboard Streamlit, nghiệm thu và bảo vệ đồ án.", "• Buổi thực hành 6:")
    add_body("Kỳ vọng: Giảm từ 15% - 25% chi phí tiền điện hàng tháng và 20% - 30% công suất đỉnh giờ cao điểm.")

    output_path = r"D:\Nam_4\RL\BAO_CAO_THUC_HANH_1_RL.docx"
    doc.save(output_path)
    print(f"Standardized document saved successfully at: {output_path}")

if __name__ == '__main__':
    create_document()
